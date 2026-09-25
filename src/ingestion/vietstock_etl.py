"""ETL: parse VietstockFinance Excel exports into the PostgreSQL schema
(infra/postgres/migrations/001_init_schema.sql - migrated 2026-07-26 off
MySQL, see docs/PROGRESS.md).

Supports the 2 per-company export formats (see docs/data_plan.md):

  - "Báo cáo tài chính" export: one file per company, ALL periods as columns,
    full statement with MS codes (Circular 200) and audit metadata.
    -> loads financial_line_items + disclosures

  - "Chỉ số tài chính" export: one file per company, ALL periods as columns,
    ratios grouped under "NHÓM CHỈ SỐ ..." headers.
    -> loads financial_ratios

Both loads are UPSERTS (INSERT ... ON CONFLICT (...) DO UPDATE SET ..., the
Postgres upsert - was MySQL's "ON DUPLICATE KEY UPDATE" before the
2026-07-26 migration) - safe to re-run on the same file, and safe to load
files with overlapping periods (the newer file's values simply overwrite the
older ones for that period).

IMPORTANT: MS codes reset per statement (e.g. code '01' means "Lợi nhuận
trước thuế" in cash_flow_indirect but "Doanh thu bán hàng..." in
income_statement). statement_type is part of the unique key for exactly
this reason - never look up a line item by ms_code alone.

Usage:
    python -m src.ingestion.vietstock_etl bao_cao_tai_chinh <file.xlsx>
    python -m src.ingestion.vietstock_etl chi_so_tai_chinh <file.xlsx>
"""
import argparse
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
from sqlalchemy import text

from src.core.db import get_session

# ---------------------------------------------------------------------------
# shared helpers
# ---------------------------------------------------------------------------

def parse_period_label(label: str) -> tuple[int, str]:
    """'Năm/2024' -> (2024,'YEAR'); 'Quý 1/2024' -> (2024,'Q1');
    '6 tháng/2024' -> (2024,'H1'); '9 tháng/2024' -> (2024,'9M')."""
    label = label.strip()
    m = re.match(r"Năm/(\d{4})", label)
    if m:
        return int(m.group(1)), "YEAR"
    m = re.match(r"Quý\s*(\d)/(\d{4})", label)
    if m:
        return int(m.group(2)), f"Q{m.group(1)}"
    m = re.match(r"6\s*tháng/(\d{4})", label)
    if m:
        return int(m.group(1)), "H1"
    m = re.match(r"9\s*tháng/(\d{4})", label)
    if m:
        return int(m.group(1)), "9M"
    raise ValueError(f"Unrecognized period label: {label!r}")


def parse_date(s) -> date | None:
    if not s:
        return None
    return datetime.strptime(str(s).strip(), "%d-%m-%Y").date()


def consolidation_from(label: str) -> str:
    return "consolidated" if "hợp nhất" in (label or "").lower() else "standalone"


def period_to_report_type(period: str) -> str:
    return "nam" if period == "YEAR" else ("ban_nien" if period in ("H1", "9M") else "quy")


# ---------------------------------------------------------------------------
# "Bao cao tai chinh" format
# ---------------------------------------------------------------------------

# Rows where column B literally contains the string "MS" mark the start of a
# new statement section (Vietstock repeats the "MS" column header each time).
STATEMENT_HEADERS = {
    "BÁO CÁO TÌNH HÌNH TÀI CHÍNH": "balance_sheet",
    "KẾT QUẢ KINH DOANH": "income_statement",
    "LƯU CHUYỂN TIỀN TỆ GIÁN TIẾP": "cash_flow_indirect",
    "LƯU CHUYỂN TIỀN TỆ TRỰC TIẾP": "cash_flow_direct",
}


def parse_bao_cao_tai_chinh(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    n_cols = ws.max_column

    tickers = [ws.cell(row=6, column=c).value for c in range(3, n_cols + 1)]
    period_labels = [ws.cell(row=7, column=c).value for c in range(3, n_cols + 1)]
    consolidations = [ws.cell(row=8, column=c).value for c in range(3, n_cols + 1)]
    audit_statuses = [ws.cell(row=9, column=c).value for c in range(3, n_cols + 1)]
    auditors = [ws.cell(row=10, column=c).value for c in range(3, n_cols + 1)]
    opinions = [ws.cell(row=11, column=c).value for c in range(3, n_cols + 1)]
    disclosure_dates = [ws.cell(row=12, column=c).value for c in range(3, n_cols + 1)]
    signed_dates = [ws.cell(row=13, column=c).value for c in range(3, n_cols + 1)]

    col_meta = []
    for i in range(len(tickers)):
        if not tickers[i] or not period_labels[i]:
            continue
        fiscal_year, period = parse_period_label(period_labels[i])
        col_meta.append({
            "col_idx": i,
            "ticker": tickers[i],
            "fiscal_year": fiscal_year,
            "period": period,
            "consolidation": consolidation_from(consolidations[i]),
            "audit_status": audit_statuses[i],
            "auditor": auditors[i],
            "audit_opinion": opinions[i],
            "disclosure_date": parse_date(disclosure_dates[i]),
            "signed_date": parse_date(signed_dates[i]),
        })

    line_items = []
    statement_type = "balance_sheet"
    for r in range(14, ws.max_row + 1):
        item_name = ws.cell(row=r, column=1).value
        ms_code = ws.cell(row=r, column=2).value
        if item_name is None:
            continue
        if ms_code == "MS":
            statement_type = STATEMENT_HEADERS.get(str(item_name).strip(), statement_type)
            continue
        if ms_code is None:
            continue  # section divider row (e.g. "TÀI SẢN"), no data of its own
        for meta in col_meta:
            value = ws.cell(row=r, column=3 + meta["col_idx"]).value
            if value is None:
                continue
            line_items.append({
                "ticker": meta["ticker"],
                "fiscal_year": meta["fiscal_year"],
                "period": meta["period"],
                "consolidation": meta["consolidation"],
                "ms_code": str(ms_code),
                "item_name": str(item_name).strip(),
                "statement_type": statement_type,
                "value": value,
            })

    disclosures = [
        {k: meta[k] for k in (
            "ticker", "fiscal_year", "period", "consolidation",
            "audit_status", "auditor", "audit_opinion",
            "disclosure_date", "signed_date",
        )}
        for meta in col_meta
    ]
    return line_items, disclosures


# ---------------------------------------------------------------------------
# "Chi so tai chinh" format
# ---------------------------------------------------------------------------

def parse_chi_so_tai_chinh(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    n_cols = ws.max_column

    tickers = [ws.cell(row=6, column=c).value for c in range(3, n_cols + 1)]
    period_labels = [ws.cell(row=7, column=c).value for c in range(3, n_cols + 1)]

    col_meta = []
    for i in range(len(tickers)):
        if not tickers[i] or not period_labels[i]:
            continue
        fiscal_year, period = parse_period_label(period_labels[i])
        col_meta.append({"col_idx": i, "ticker": tickers[i], "fiscal_year": fiscal_year, "period": period})

    ratios = []
    ratio_group = None
    for r in range(8, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        unit = ws.cell(row=r, column=2).value
        if name is None:
            continue
        if str(name).startswith("NHÓM") and unit is None:
            ratio_group = str(name).strip()
            continue
        for meta in col_meta:
            value = ws.cell(row=r, column=3 + meta["col_idx"]).value
            if value is None:
                continue
            ratios.append({
                "ticker": meta["ticker"],
                "fiscal_year": meta["fiscal_year"],
                "period": meta["period"],
                "ratio_group": ratio_group,
                "ratio_name": str(name).strip(),
                "value": value,
                "unit": unit,
            })
    return ratios


# ---------------------------------------------------------------------------
# DB load (upsert)
# ---------------------------------------------------------------------------

UPSERT_LINE_ITEM = text("""
    INSERT INTO financial_line_items
        (ticker, fiscal_year, period, ms_code, item_name, statement_type, consolidation, value)
    VALUES
        (:ticker, :fiscal_year, :period, :ms_code, :item_name, :statement_type, :consolidation, :value)
    ON CONFLICT (ticker, fiscal_year, period, statement_type, ms_code, consolidation)
    DO UPDATE SET item_name = EXCLUDED.item_name, value = EXCLUDED.value
""")

UPSERT_DISCLOSURE = text("""
    INSERT INTO disclosures
        (ticker, fiscal_year, period, consolidation, report_type,
         audit_status, auditor, audit_opinion, disclosure_date, signed_date)
    VALUES
        (:ticker, :fiscal_year, :period, :consolidation, :report_type,
         :audit_status, :auditor, :audit_opinion, :disclosure_date, :signed_date)
    ON CONFLICT (ticker, fiscal_year, period, consolidation)
    DO UPDATE SET
        audit_status = EXCLUDED.audit_status, auditor = EXCLUDED.auditor,
        audit_opinion = EXCLUDED.audit_opinion, disclosure_date = EXCLUDED.disclosure_date,
        signed_date = EXCLUDED.signed_date
""")

UPSERT_RATIO = text("""
    INSERT INTO financial_ratios
        (ticker, fiscal_year, period, ratio_group, ratio_name, value, unit)
    VALUES
        (:ticker, :fiscal_year, :period, :ratio_group, :ratio_name, :value, :unit)
    ON CONFLICT (ticker, fiscal_year, period, ratio_name)
    DO UPDATE SET value = EXCLUDED.value, unit = EXCLUDED.unit
""")

INSERT_LOG = text("""
    INSERT INTO ingestion_log (file_name, file_type, ticker, row_count, notes)
    VALUES (:file_name, :file_type, :ticker, :row_count, :notes)
""")

CHECK_TICKER = text("SELECT 1 FROM companies WHERE ticker = :ticker")

# Auto-register any ticker the ETL encounters but isn't in `companies` yet,
# instead of silently skipping the file (2026-07-21: VNZ/VCB/VIB/AAN were
# found silently dropped this way - files sat in data/raw/vietstock_exports/
# for days with zero indication anything was wrong beyond a console print
# easy to miss when batch-processing). Placeholder metadata mirrors
# scripts/seed_companies.py's convention for top-200 candidates
# (sector="Chưa phân loại") - update companies.company_name/exchange/sector
# by hand (or re-run seed_companies.py) once you decide the ticker belongs
# in the thesis's evaluated universe. `ON CONFLICT (ticker) DO NOTHING`
# is a no-op upsert so this never clobbers real metadata seeded elsewhere
# (Postgres equivalent of MySQL's old `ON DUPLICATE KEY UPDATE ticker=ticker`
# self-assignment trick - DO NOTHING is the direct, non-hacky way to say it).
REGISTER_PLACEHOLDER_TICKER = text("""
    INSERT INTO companies (ticker, company_name, exchange, sector)
    VALUES (:ticker, :ticker, 'UNKNOWN', 'Chưa phân loại')
    ON CONFLICT (ticker) DO NOTHING
""")


def _ticker_in_companies(session, ticker: str) -> bool:
    return session.execute(CHECK_TICKER, {"ticker": ticker}).first() is not None


def _ensure_ticker_registered(session, ticker: str, path_name: str) -> None:
    if not _ticker_in_companies(session, ticker):
        session.execute(REGISTER_PLACEHOLDER_TICKER, {"ticker": ticker})
        print(f"{path_name}: ticker '{ticker}' not in companies table - "
              f"AUTO-REGISTERED as placeholder (company_name=ticker, sector='Chưa phân loại'). "
              f"Fix real company_name/exchange/sector by hand if this ticker stays in scope.")


def load_bao_cao_tai_chinh(path: Path):
    line_items, disclosures = parse_bao_cao_tai_chinh(path)
    ticker = disclosures[0]["ticker"] if disclosures else None

    with get_session() as session:
        if ticker:
            _ensure_ticker_registered(session, ticker, path.name)
        for row in line_items:
            session.execute(UPSERT_LINE_ITEM, row)
        for row in disclosures:
            session.execute(UPSERT_DISCLOSURE, {**row, "report_type": period_to_report_type(row["period"])})
        session.execute(INSERT_LOG, {
            "file_name": path.name,
            "file_type": "bao_cao_tai_chinh",
            "ticker": ticker,
            "row_count": len(line_items),
            "notes": f"{len(disclosures)} periods",
        })
    print(f"{path.name}: loaded {len(line_items)} line items across {len(disclosures)} periods")


def load_chi_so_tai_chinh(path: Path):
    ratios = parse_chi_so_tai_chinh(path)
    ticker = ratios[0]["ticker"] if ratios else None

    with get_session() as session:
        if ticker:
            _ensure_ticker_registered(session, ticker, path.name)
        for row in ratios:
            session.execute(UPSERT_RATIO, row)
        session.execute(INSERT_LOG, {
            "file_name": path.name,
            "file_type": "chi_so_tai_chinh",
            "ticker": ticker,
            "row_count": len(ratios),
            "notes": None,
        })
    print(f"{path.name}: loaded {len(ratios)} ratio values")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("kind", choices=["bao_cao_tai_chinh", "chi_so_tai_chinh"])
    ap.add_argument("file", type=str)
    args = ap.parse_args()

    if args.kind == "bao_cao_tai_chinh":
        load_bao_cao_tai_chinh(Path(args.file))
    else:
        load_chi_so_tai_chinh(Path(args.file))
