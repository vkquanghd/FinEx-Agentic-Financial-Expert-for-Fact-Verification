"""ETL: parse VietstockFinance "Du lieu vi mo" (macro/national-economy)
Excel exports into macro_indicators.

Each file covers ONE macro topic (title in row 3) across ALL periods of a
single granularity - quarterly (e.g. GDP, CPI-by-quarter... note CPI/PMI/
retail/IIP are actually monthly), monthly, or annual (population,
industry-breakdown files). Header row (containing "Chỉ tiêu" in column A)
marks where period columns start; everything below is one row per
indicator, one column per period.

Rows with no parseable numeric value in ANY period column are skipped -
this silently and correctly drops section-header rows (e.g. a sector name
repeated as a blank divider before its real data row) and Vietstock's
footer/contact-info rows, without needing to hardcode row positions (which
differ per file).

Some exports (interest rates, FX rate, open-market-operation rates) report
one column PER CALENDAR DAY (dd/mm/yyyy) instead of quý/tháng/năm - these
get period_type='DAILY', period='D<MMDD>'.

A few Vietstock "macro" exports are NOT in the indicator-x-period-column
shape at all (e.g. "KẾT QUẢ ĐẤU THẦU THỊ TRƯỜNG MỞ" is one row per auction
event, columns are transaction fields, not periods) - these don't fit
macro_indicators and are skipped with an explicit message rather than
force-fit or crash the whole batch.

Load is an UPSERT - safe to re-run.

Usage:
    python -m src.ingestion.macro_etl data/raw/vietstock_exports/macro/*.xlsx
"""
import argparse
import glob
import re
from pathlib import Path

import openpyxl
from sqlalchemy import text

from src.core.db import get_session

# Title (row 3, first non-empty cell) -> stable category code. Falls back to
# a slugified version of the title for any file title not listed here, so
# new macro exports don't crash the loader, just get an auto-generated code.
TITLE_TO_CATEGORY = {
    "GDP DANH NGHĨA": "GDP_NOMINAL",
    "CƠ CẤU GDP DANH NGHĨA": "GDP_NOMINAL_STRUCTURE",
    "GDP THỰC": "GDP_REAL",
    "TĂNG TRƯỞNG GDP THỰC": "GDP_REAL_GROWTH",
    "CHỈ SỐ GIÁ TIÊU DÙNG (CPI)": "CPI",
    "CHỈ SỐ GIÁ SẢN XUẤT (PPI)": "PPI",
    "CHỈ SỐ GIÁ NGUYÊN NHIÊN VẬT LIỆU DÙNG CHO SẢN XUẤT (IPI)": "IPI",
    "CHỈ SỐ GIÁ XUẤT KHẨU HÀNG HÓA (XPI)": "XPI",
    "CHỈ SỐ GIÁ NHẬP KHẨU HÀNG HÓA (MPI)": "MPI",
    "LAO ĐỘNG": "LABOR",
    "DÂN SỐ": "POPULATION",
    "TỶ LỆ THẤT NGHIỆP": "UNEMPLOYMENT",
    "BÁN LẺ": "RETAIL",
    "CHỈ SỐ NHÀ QUẢN TRỊ MUA HÀNG (PMI)": "PMI",
    "CHỈ SỐ SẢN XUẤT CÔNG NGHIỆP (IIP)": "IIP",
    "KHỐI LƯỢNG SẢN XUẤT CÔNG NGHIỆP": "INDUSTRIAL_VOLUME",
    "CHỈ SỐ SẢN XUẤT CÔNG NGHIỆP PHÂN THEO NGÀNH CÔNG NGHIỆP": "IIP_BY_INDUSTRY",
    "CHỈ SỐ TIÊU THỤ NGÀNH CÔNG NGHIỆP CHẾ BIẾN, CHẾ TẠO PHÂN THEO NGÀNH CÔNG NGHIỆP": "CONSUMPTION_BY_INDUSTRY",
    "CHỈ SỐ TỒN KHO NGÀNH CÔNG NGHIỆP CHẾ BIẾN, CHẾ TẠO TẠI THỜI ĐIỂM 31/12 HÀNG NĂM PHÂN THEO NGÀNH CÔNG NGHIỆP": "INVENTORY_BY_INDUSTRY",
    "VAY NỢ CHÍNH PHỦ": "GOV_DEBT",
    "VAY NỢ CHÍNH PHỦ BẢO LÃNH": "GOV_GUARANTEED_DEBT",
    "VAY NỢ CHÍNH QUYỀN ĐỊA PHƯƠNG": "LOCAL_GOV_DEBT",
    "VAY NỢ NGOÀI QUỐC GIA": "NATIONAL_EXTERNAL_DEBT",
    "SỐ LƯỢNG DOANH NGHIỆP ĐĂNG KÝ KINH DOANH VÀ GIẢI THỂ": "BUSINESS_REGISTRATION",
    "ĐẦU TƯ PHÁT TRIỂN TOÀN XÃ HỘI": "SOCIAL_INVESTMENT",
    "ĐẦU TƯ NGÂN SÁCH NHÀ NƯỚC": "STATE_BUDGET_INVESTMENT",
    "DỰ TRỮ NGOẠI HỐI": "FX_RESERVES",
    "TÍN DỤNG": "CREDIT",
    "HUY ĐỘNG": "DEPOSITS",
    "KẾT QUẢ ĐẤU THẦU THỊ TRƯỜNG MỞ": "OMO_AUCTION_RESULTS",  # different shape, see module docstring - will be SKIPPED
    "THỊ TRƯỜNG MỞ": "OMO_RATES",
    "TỶ GIÁ TRUNG TÂM USD/VND": "USDVND_CENTRAL_RATE",
    "LÃI SUẤT LIÊN NGÂN HÀNG VNIBOR": "VNIBOR",
    "LÃI SUẤT TIỀN GỬI NGÂN HÀNG": "DEPOSIT_RATES",
    "NGÀNH": "FDI_BY_INDUSTRY",
    "XUẤT KHẨU PHÂN THEO NƯỚC, KHỐI NƯỚC VÀ VÙNG LÃNH THỔ CHỦ YẾU": "EXPORTS_BY_COUNTRY",
    "NHẬP KHẨU PHÂN THEO NƯỚC, KHỐI NƯỚC VÀ VÙNG LÃNH THỔ CHỦ YẾU": "IMPORTS_BY_COUNTRY",
    "XUẤT KHẨU THEO HÀNG HÓA": "EXPORTS_BY_COMMODITY",
    "NHẬP KHẨU THEO HÀNG HÓA": "IMPORTS_BY_COMMODITY",
    "MỘT SỐ SẢN PHẨM CÔNG NGHIỆP CHỦ YẾU": "KEY_INDUSTRIAL_PRODUCTS",
    "MỘT SỐ SẢN PHẨM CÔNG NGHIỆP CHỦ YẾU PHÂN THEO LOẠI HÌNH KINH TẾ": "KEY_INDUSTRIAL_PRODUCTS_BY_OWNERSHIP",
    "BẢNG CÂN ĐỐI NĂNG LƯỢNG VIỆT NAM": "ENERGY_BALANCE",
    "DIỆN TÍCH SÀN XÂY DỰNG NHÀ TỰ XÂY, TỰ Ở HOÀN THÀNH TRONG NĂM CỦA HỘ DÂN CƯ PHÂN THEO LOẠI NHÀ VÀ THEO MỨC ĐỘ KIÊN CỐ": "SELF_BUILT_FLOOR_AREA",
    "DIỆN TÍCH SÀN XÂY DỰNG NHÀ Ở HOÀN THÀNH TRONG NĂM PHÂN THEO LOẠI NHÀ": "COMPLETED_FLOOR_AREA",
    "CÁN CÂN THANH TOÁN QUỐC TẾ": "BALANCE_OF_PAYMENTS",
    "THU NGÂN SÁCH NHÀ NƯỚC": "STATE_BUDGET_REVENUE",
    "CHI NGÂN SÁCH NHÀ NƯỚC": "STATE_BUDGET_EXPENDITURE",
}

QUARTER_RE = re.compile(r"Qu[yý]\s*(\d)\s*/\s*(\d{4})")
MONTH_RE = re.compile(r"Th[aá]ng\s*(\d{1,2})\s*/\s*(\d{4})")
YEAR_RE = re.compile(r"^(\d{4})$")
DAILY_RE = re.compile(r"^(\d{2})/(\d{2})/(\d{4})$")  # dd/mm/yyyy - FX/interest-rate exports
# Government budget exports report YTD-cumulative figures at 3/6/9/12 months
# (e.g. "9 Tháng/2021" = cumulative through 9 months of 2021) - number comes
# BEFORE "Tháng", unlike the calendar-month header ("Tháng 7/2021") above.
# This is a genuinely different thing (year-to-date cumulative, not a point
# value for that single month) so it gets its own period_type rather than
# being silently relabeled as a quarter/half-year point value.
YTD_MONTH_RE = re.compile(r"^(\d{1,2})\s*Th[aá]ng\s*/\s*(\d{4})$")


def _slugify(title: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", title.upper()).strip("_")[:50]


def _parse_period_header(label: str):
    """Return (period_type, fiscal_year, period) or None if not a period column."""
    if label is None:
        return None
    label = str(label).strip()
    m = QUARTER_RE.match(label)
    if m:
        return "QUARTER", int(m.group(2)), f"Q{m.group(1)}"
    m = YTD_MONTH_RE.match(label)
    if m:
        return "YTD", int(m.group(2)), f"YTD{m.group(1)}"
    m = MONTH_RE.match(label)
    if m:
        return "MONTH", int(m.group(2)), f"M{m.group(1)}"
    m = YEAR_RE.match(label)
    if m:
        return "YEAR", int(m.group(1)), "YEAR"
    m = DAILY_RE.match(label)
    if m:
        dd, mm, yyyy = m.groups()
        return "DAILY", int(yyyy), f"D{mm}{dd}"
    return None


def _find_title(ws) -> str:
    for r in range(1, 6):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v and str(v).strip():
                return str(v).strip()
    return "UNKNOWN"


def _find_header_row(ws) -> int:
    for r in range(1, min(ws.max_row, 15) + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == "Chỉ tiêu":
            return r
    raise ValueError("Could not find header row (column A == 'Chỉ tiêu')")


def parse_macro_file(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    title = _find_title(ws)
    category = TITLE_TO_CATEGORY.get(title, _slugify(title))
    header_row = _find_header_row(ws)

    period_cols = []  # list of (col_idx, period_type, fiscal_year, period)
    for c in range(3, ws.max_column + 1):
        parsed = _parse_period_header(ws.cell(row=header_row, column=c).value)
        if parsed:
            period_cols.append((c, *parsed))

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        unit = ws.cell(row=r, column=2).value
        if name is None or not str(name).strip():
            continue
        name = str(name).strip()

        row_values = []
        for c, period_type, fiscal_year, period in period_cols:
            raw = ws.cell(row=r, column=c).value
            if raw is None or isinstance(raw, str):
                continue
            row_values.append((period_type, fiscal_year, period, float(raw)))

        if not row_values:
            # No numeric data on this row at all - either a section-header
            # divider (e.g. sector name repeated with blank cells before its
            # real data row) or Vietstock footer/contact-info text. Skip.
            continue

        for period_type, fiscal_year, period, value in row_values:
            rows.append({
                "category": category,
                "indicator_name": name,
                "unit": str(unit).strip() if unit else None,
                "period_type": period_type,
                "fiscal_year": fiscal_year,
                "period": period,
                "value": value,
            })

    return category, title, rows


UPSERT_MACRO = text("""
    INSERT INTO macro_indicators
        (category, indicator_name, unit, period_type, fiscal_year, period, value, source_file)
    VALUES
        (:category, :indicator_name, :unit, :period_type, :fiscal_year, :period, :value, :source_file)
    ON CONFLICT (category, indicator_name, period_type, fiscal_year, period)
    DO UPDATE SET
        unit = EXCLUDED.unit, value = EXCLUDED.value, source_file = EXCLUDED.source_file
""")

INSERT_LOG = text("""
    INSERT INTO ingestion_log (file_name, file_type, ticker, row_count, notes)
    VALUES (:file_name, :file_type, NULL, :row_count, :notes)
""")


def load_macro_file(path: Path):
    try:
        category, title, rows = parse_macro_file(path)
    except ValueError as e:
        # E.g. no "Chỉ tiêu" header row found - this file isn't in the
        # indicator-x-period-column shape (see module docstring: some
        # Vietstock "macro" exports are transaction-level tables instead).
        # Skip it and keep going rather than aborting the whole batch.
        print(f"{path.name}: SKIPPED - {e}")
        return

    if not rows:
        print(f"{path.name}: SKIPPED - no parseable data rows found (title='{title}')")
        return

    with get_session() as session:
        for row in rows:
            session.execute(UPSERT_MACRO, {**row, "source_file": path.name})
        session.execute(INSERT_LOG, {
            "file_name": path.name,
            "file_type": "du_lieu_vi_mo",
            "row_count": len(rows),
            "notes": f"category={category} ({title})",
        })
    n_indicators = len({r["indicator_name"] for r in rows})
    print(f"{path.name}: loaded {len(rows)} values, category={category}, {n_indicators} indicators, title='{title}'")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+", help="One or more macro export .xlsx files (globs expanded by shell or here)")
    args = ap.parse_args()

    paths = []
    for pattern in args.files:
        matched = glob.glob(pattern)
        paths.extend(matched if matched else [pattern])

    for p in paths:
        try:
            load_macro_file(Path(p))
        except Exception as e:  # e.g. corrupted/incomplete .xlsx download
            print(f"{Path(p).name}: SKIPPED - unexpected error: {e}")
