"""Data-coverage audit (2026-08-06) - answers "công ty nào có ĐỦ file gì?"

Two halves:
  1. FILE half (no DB needed): scans data/raw/vietstock_exports/*.xlsx
     headers (row 6 = ticker, row 7 = period label) to report, per ticker,
     which bao-cao/chi-so periods exist ON DISK.
  2. DB half (needs the local Postgres up): per-ticker coverage across the 4
     modalities (financial_line_items periods, financial_ratios,
     company_price_snapshot, disclosure_chunks doc counts) - i.e. what is
     actually QUERYABLE by the app, which is what the Chat demo cares about.

Writes docs/data_coverage_report.md so the coverage matrix can be cited in
the thesis (data section) and used to pick the "trọng tâm" ticker set the
demo should advertise as fully supported.

Usage: python3 -m scripts.audit_data_coverage [--files-only]
"""
import argparse
import os
import re
import sys
import warnings
from collections import defaultdict
from pathlib import Path

warnings.filterwarnings("ignore")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

EXPORT_DIR = ROOT / "data/raw/vietstock_exports"
OUT_MD = ROOT / "docs/data_coverage_report.md"

CORE_17 = ["VNM", "SAB", "MSN", "KDC", "VIC", "VHM", "NVL", "KDH",
           "HPG", "GAS", "DPM", "DCM", "MWG", "PNJ", "DGW", "FRT", "VNZ"]


def scan_export_files():
    from openpyxl import load_workbook
    cov = defaultdict(lambda: {"bao_cao": set(), "chi_so": set()})
    n_empty = n_bad = 0
    files = sorted(EXPORT_DIR.glob("*.xlsx"))
    for f in files:
        if f.stat().st_size == 0:
            n_empty += 1
            continue
        kind = "bao_cao" if "Bao-cao" in f.name else "chi_so"
        try:
            wb = load_workbook(f, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(min_row=6, max_row=7, min_col=3,
                                     max_col=ws.max_column or 3, values_only=True))
            for t, p in zip(rows[0], rows[1]):
                if t and p:
                    cov[str(t).strip()][kind].add(str(p).strip())
            wb.close()
        except Exception:
            n_bad += 1
    return cov, len(files), n_empty, n_bad


def year_range(periods):
    ys = sorted({int(m.group(1)) for p in periods
                 for m in [re.search(r"(\d{4})", p)] if m})
    return f"{min(ys)}-{max(ys)}" if ys else "-"


def db_coverage():
    import pandas as pd
    from src.core.db import engine
    fli = pd.read_sql(
        "SELECT ticker, COUNT(DISTINCT (fiscal_year, period)) AS n_periods, "
        "MIN(fiscal_year) AS y0, MAX(fiscal_year) AS y1, "
        "COUNT(DISTINCT period) AS period_kinds "
        "FROM financial_line_items GROUP BY ticker", engine)
    ratios = pd.read_sql(
        "SELECT ticker, COUNT(DISTINCT (fiscal_year, period)) AS n_ratio_periods "
        "FROM financial_ratios GROUP BY ticker", engine)
    prices = pd.read_sql(
        "SELECT ticker, COUNT(*) AS n_price_rows FROM company_price_snapshot "
        "GROUP BY ticker", engine)
    chunks = pd.read_sql(
        "SELECT ticker, COUNT(*) AS n_chunks, COUNT(DISTINCT doctype) AS n_doctypes "
        "FROM disclosure_chunks GROUP BY ticker", engine)
    df = fli.merge(ratios, on="ticker", how="outer") \
            .merge(prices, on="ticker", how="outer") \
            .merge(chunks, on="ticker", how="outer").fillna(0)
    for c in df.columns:
        if c != "ticker":
            df[c] = df[c].astype(int)
    return df.sort_values("n_periods", ascending=False)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--files-only", action="store_true",
                    help="skip the DB half (no Postgres needed)")
    args = ap.parse_args()

    lines = ["# Data coverage report", "",
             f"Sinh bởi scripts/audit_data_coverage.py.", ""]

    cov, n_files, n_empty, n_bad = scan_export_files()
    lines += ["## 1. File Vietstock export trên đĩa", "",
              f"- Tổng file: {n_files} (rỗng/bỏ qua: {n_empty}, lỗi đọc: {n_bad})",
              f"- Số ticker xuất hiện trong export: {len(cov)}", "",
              "| Ticker | Kỳ BCTC | Năm BCTC | Kỳ chỉ số | Năm chỉ số | Ghi chú |",
              "|---|---|---|---|---|---|"]
    for t in CORE_17:
        c = cov.get(t)
        if not c or (not c["bao_cao"] and not c["chi_so"]):
            lines.append(f"| {t} | 0 | - | 0 | - | KHÔNG có file export |")
        else:
            note = "core" if t != "VNZ" else "ad-hoc (Vina Game)"
            lines.append(f"| {t} | {len(c['bao_cao'])} | {year_range(c['bao_cao'])} "
                         f"| {len(c['chi_so'])} | {year_range(c['chi_so'])} | {note} |")
    others = sorted(set(cov) - set(CORE_17))
    rich = [t for t in others if len(cov[t]["bao_cao"]) >= 10]
    lines += ["", f"Ticker NGOÀI core có >=10 kỳ BCTC trên đĩa ({len(rich)}): "
              + " ".join(rich), ""]

    if not args.files_only:
        try:
            df = db_coverage()
            lines += ["## 2. Coverage trong DB (cái app thật sự tra được)", "",
                      "| Ticker | Kỳ line-items | Năm | Kỳ ratios | Giá | Chunks RAG | Doctypes |",
                      "|---|---|---|---|---|---|---|"]
            dfc = df[df.ticker.isin(CORE_17)].set_index("ticker")
            for t in CORE_17:
                if t in dfc.index:
                    r = dfc.loc[t]
                    yr = f"{r.y0}-{r.y1}" if r.n_periods else "-"
                    lines.append(f"| {t} | {r.n_periods} | {yr} | {r.n_ratio_periods} "
                                 f"| {r.n_price_rows} | {r.n_chunks} | {r.n_doctypes} |")
                else:
                    lines.append(f"| {t} | 0 | - | 0 | 0 | 0 | 0 |")
            full = df[(df.n_periods > 0) & (df.n_chunks > 0) & (df.n_price_rows > 0)]
            lines += ["", f"Ticker đủ CẢ 3 (line-items + RAG + giá): {len(full)} — "
                      + " ".join(sorted(full.ticker)), ""]
        except Exception as e:
            lines += ["## 2. Coverage trong DB", "",
                      f"(Bỏ qua - không kết nối được DB: {type(e).__name__}: {e})", ""]

    OUT_MD.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    print(f"\nĐã ghi {OUT_MD.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
